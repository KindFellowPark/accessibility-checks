"""checks/docx_structure.py: heading order, visual headings, table header rows, language and
title. Each fixture is a real .docx built with python-docx."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from checks import docx_structure as gate


def build(path, body, title="Fixture"):
    d = docx.Document()
    d.core_properties.title = title
    d.add_paragraph("Real heading", style="Heading 1")
    body(d)
    d.save(path)
    return path


def visual(path):
    return [f for f in gate.check_docx(path) if f.startswith("visual heading")]


def _header_row(table):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(el)


# ---- visual headings --------------------------------------------------------

def test_all_bold_short_paragraph_over_body_is_a_visual_heading(tmp_path):
    def body(d):
        d.add_paragraph().add_run("Next steps").bold = True
        d.add_paragraph("Body text under the bold line.")
    path = build(tmp_path / "fake.docx", body)
    assert len(visual(path)) == 1
    assert "Next steps" in visual(path)[0]


def test_real_heading_style_passes(tmp_path):
    def body(d):
        d.add_paragraph("Next steps", style="Heading 2")
        d.add_paragraph("Body text.")
    assert visual(build(tmp_path / "styled.docx", body)) == []


def test_bold_label_inside_a_sentence_passes(tmp_path):
    def body(d):
        p = d.add_paragraph()
        p.add_run("Deadline: ").bold = True
        p.add_run("Friday at noon.")
        d.add_paragraph("Body text.")
    assert visual(build(tmp_path / "partial.docx", body)) == []


def test_bold_closing_line_passes(tmp_path):
    def body(d):
        d.add_paragraph("Body text.")
        d.add_paragraph().add_run("Thank you.").bold = True
    assert visual(build(tmp_path / "closing.docx", body)) == []


def test_bold_table_cell_passes(tmp_path):
    def body(d):
        t = d.add_table(rows=2, cols=1)
        t.rows[0].cells[0].paragraphs[0].add_run("Column").bold = True
        t.rows[1].cells[0].text = "value"
        d.add_paragraph("Body text.")
    assert visual(build(tmp_path / "table.docx", body)) == []


def test_long_bold_paragraph_is_emphasis_not_a_heading(tmp_path):
    def body(d):
        d.add_paragraph().add_run("This is a long bold sentence that runs well past the "
                                  "twelve word ceiling and is emphasis, not a heading.").bold = True
        d.add_paragraph("Body text.")
    assert visual(build(tmp_path / "long.docx", body)) == []


def test_bold_paragraph_style_counts_as_bold(tmp_path):
    def body(d):
        st = d.styles.add_style("BoldBody", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        st.font.bold = True
        d.add_paragraph("Why it matters", style="BoldBody")
        d.add_paragraph("Body text.")
    assert len(visual(build(tmp_path / "boldstyle.docx", body))) == 1


# ---- heading order, tables, language, title ---------------------------------

def test_heading_skip_fails(tmp_path):
    def body(d):
        d.add_paragraph("Too deep", style="Heading 3")
    failures = gate.check_docx(build(tmp_path / "skip.docx", body))
    assert any("skip from Heading 1 to Heading 3" in f for f in failures)


def test_table_without_header_row_fails_and_with_one_passes(tmp_path):
    def no_header(d):
        d.add_table(rows=2, cols=2)
    def with_header(d):
        _header_row(d.add_table(rows=2, cols=2))
    assert any(f.startswith("table header")
               for f in gate.check_docx(build(tmp_path / "t1.docx", no_header)))
    assert gate.check_docx(build(tmp_path / "t2.docx", with_header)) == []


def test_missing_title_fails(tmp_path):
    failures = gate.check_docx(build(tmp_path / "notitle.docx", lambda d: None, title=""))
    assert any(f.startswith("title") for f in failures)


def test_clean_document_passes_with_exit_0(tmp_path):
    path = build(tmp_path / "clean.docx", lambda d: d.add_paragraph("Body text."))
    assert gate.main(["x", str(path)]) == 0


def test_no_argument_exits_2():
    assert gate.main(["x"]) == 2
