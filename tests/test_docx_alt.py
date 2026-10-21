"""checks/docx_alt.py: real alt text passes; missing, empty, a file name or a tool's
placeholder fails; an image marked decorative needs none. Each case builds a real .docx."""
import re
import shutil
import zipfile

import pytest
from docx import Document
from docx.shared import Inches

from checks import docx_alt as gate
from tests._png import write_png


def _set_alt(path, alt):
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                data = re.sub(rb' descr="[^"]*"', b"", data)
                if alt is not None:
                    data = re.sub(rb"(<wp:docPr[^>]*?)(/?)>",
                                  rb'\1 descr="' + alt.encode() + rb'"\2>', data)
            dst.writestr(item, data)
    shutil.move(tmp, path)


@pytest.fixture
def make_doc(tmp_path):
    img = write_png(tmp_path / "dot.png")

    def build(name, alt):
        d = Document()
        d.add_paragraph("Body text.")
        d.add_picture(str(img), width=Inches(1))
        path = tmp_path / name
        d.save(str(path))
        _set_alt(path, alt)
        return path
    return build


def test_real_alt_text_passes(make_doc):
    path = make_doc("real.docx", "Students crossing the stage at commencement")
    assert gate.main(["x", str(path)]) == 0
    assert gate.missing_alt(path) == []


def test_no_alt_attribute_fails(make_doc):
    path = make_doc("none.docx", None)
    assert gate.main(["x", str(path)]) == 1
    assert gate.missing_alt(path)[0][3] == "no alt text"


def test_whitespace_alt_fails(make_doc):
    assert gate.main(["x", str(make_doc("empty.docx", "   "))]) == 1


def test_file_name_as_alt_fails(make_doc):
    path = make_doc("filename.docx", "dot.png")
    assert gate.main(["x", str(path)]) == 1
    assert gate.missing_alt(path)[0][3] == "its file name"


def test_placeholder_alt_fails(make_doc):
    path = make_doc("placeholder.docx", "Picture 3")
    assert gate.main(["x", str(path)]) == 1
    assert gate.missing_alt(path)[0][3] == "a placeholder a tool wrote"


def test_real_description_starting_with_a_placeholder_word_passes(make_doc):
    path = make_doc("close.docx", "Chart of enrollment by term, 2021 to 2026")
    assert gate.main(["x", str(path)]) == 0


def test_image_count(make_doc):
    assert gate.count_images(make_doc("count.docx", "A navy square")) == 1


def test_no_argument_is_a_usage_error():
    assert gate.main(["x"]) == 2
