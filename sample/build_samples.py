"""Builds the sample files: two images and the before and after Word memos.

The HTML pages are written by hand and live beside this script. The memos are built here
with python-docx so the defects in the "before" memo are exact and repeatable.

    python3 sample/build_samples.py
"""
import re
import shutil
import struct
import zipfile
import zlib
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
NAVY, SLATE, GOLD, INK = (0x1F, 0x3A, 0x5F), (0x4A, 0x55, 0x68), (0xB7, 0x79, 0x1F), (0x1A, 0x20, 0x2C)
WHITE, GOLD_TINT = (0xFF, 0xFF, 0xFF), (0xFD, 0xF6, 0xE7)


# ── images, standard library only ───────────────────────────────────────────
def write_png(path, width, height, pixel):
    """A real RGB PNG. `pixel(x, y)` returns an (r, g, b) tuple."""
    raw = b"".join(b"\x00" + bytes(c for x in range(width) for c in pixel(x, y))
                   for y in range(height))

    def chunk(kind, data):
        return (struct.pack(">I", len(data)) + kind + data
                + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF))

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))
    Path(path).write_bytes(png)


def quad_header(path):
    """A flat illustration: navy sky, gold sun, slate buildings, a white path."""
    w, h = 720, 240

    def px(x, y):
        if (x - 560) ** 2 + (y - 80) ** 2 < 38 ** 2:
            return GOLD
        buildings = [(60, 200, 110), (230, 330, 70), (360, 470, 125), (500, 660, 95)]
        for x0, x1, top in buildings:
            if x0 <= x < x1 and y >= top:
                if (x - x0) % 30 in range(10, 20) and (y - top) % 34 in range(12, 24):
                    return GOLD_TINT
                return SLATE
        if y >= 205:
            return WHITE if abs(x - 360) < (y - 200) * 3 else (0xD5, 0xDA, 0xE1)
        return NAVY
    write_png(path, w, h, px)


def rsvp_chart(path):
    """A bar chart of RSVPs by week: 40, 65, 90, 120, on a white field with a slate axis."""
    w, h = 480, 260
    values = [40, 65, 90, 120]

    def px(x, y):
        if y == 230 and 40 <= x <= 450:
            return SLATE
        if x == 40 and 20 <= y <= 230:
            return SLATE
        for i, v in enumerate(values):
            x0 = 70 + i * 95
            if x0 <= x < x0 + 60 and 230 - int(v * 1.6) <= y < 230:
                return NAVY if i < 3 else GOLD
        return WHITE
    write_png(path, w, h, px)


# ── Word memos ───────────────────────────────────────────────────────────────
def _set_metadata(doc, title):
    cp = doc.core_properties
    cp.title = title
    cp.author = "Jordan Lee"
    cp.last_modified_by = "Jordan Lee"
    cp.subject = "Fall Community Open House" if title else ""
    cp.comments = ""
    cp.keywords = ""
    cp.category = ""
    cp.revision = 1
    cp.created = cp.modified = datetime(2026, 10, 1, 9, 0, 0)


def _mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _postprocess(path, strip_alt):
    """Drop the blank-template thumbnail python-docx ships (a stale preview that also carries
    image metadata) and, for the "before" memo, make sure no alt text survives."""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename == "docProps/thumbnail.jpeg":
                continue
            data = src.read(item.filename)
            if item.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship [^>]*thumbnail\.jpeg"/>', b"", data)
            if strip_alt and item.filename == "word/document.xml":
                data = re.sub(rb' descr="[^"]*"', b"", data)
            dst.writestr(item, data)
    shutil.move(tmp, path)


def _set_alt(doc, picture_run_paragraph, alt):
    for doc_pr in picture_run_paragraph._p.iter(qn("wp:docPr")):
        doc_pr.set("descr", alt)


def _body(doc, text, color=INK):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(*color)
    return p


def build_memo(path, fixed):
    doc = Document()
    # Every color in the memo comes from the declared pairings, headings included.
    for name in ("Title", "Heading 1"):
        doc.styles[name].font.color.rgb = RGBColor(*NAVY)
    for bdr in doc.styles["Title"].element.iter(qn("w:pBdr")):
        bdr.getparent().remove(bdr)  # the template's blue rule under the title
    _set_metadata(doc, "Fall Community Open House plan" if fixed else "")

    doc.add_paragraph("Fall Community Open House plan", style="Title")
    _body(doc, "To: Dr. Helen Marsh, President", SLATE)
    _body(doc, "From: Dana Whitcombe, Executive Director, Office of Communications", SLATE)
    _body(doc, "Prepared by: Jordan Lee, Executive Assistant", SLATE)
    _body(doc, "Date: October 1, 2026", SLATE)

    doc.add_paragraph("Purpose", style="Heading 1")
    _body(doc, "This memo sets out the plan for the Fall Community Open House on Saturday, "
               "November 14, 2026, and asks for your approval of the schedule below.")

    if fixed:
        doc.add_paragraph("Schedule", style="Heading 1")
    else:
        # Defect: looks like a heading, is only bold text.
        p = doc.add_paragraph()
        r = p.add_run("Schedule")
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(*NAVY)
    _body(doc, "Doors open at 9:00 a.m. Each department hosts a table in the Student Center, "
               "and campus tours leave from the main entrance every half hour.")

    rows = [("Time", "Event", "Location"),
            ("9:00 a.m.", "Welcome from the President", "Student Center"),
            ("9:30 a.m.", "Department tables open", "Student Center"),
            ("10:00 a.m.", "Campus tours begin", "Main entrance"),
            ("12:00 p.m.", "Closing and lunch", "Library lawn")]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = text
            if i == 0:
                cell.paragraphs[0].runs[0].bold = True
    if fixed:
        _mark_header_row(table.rows[0])  # Defect when missing: no header row.

    doc.add_paragraph("RSVPs so far", style="Heading 1")
    chart = HERE / "assets" / "rsvp_chart.png"
    doc.add_picture(str(chart), width=Inches(4))
    if fixed:
        _set_alt(doc, doc.paragraphs[-1],
                 "Bar chart of RSVPs by week: 40 in week one, 65 in week two, 90 in week "
                 "three and 120 in week four.")
    _body(doc, "RSVPs have tripled since the first week, to 120. We expect about 200 guests.")

    deadline = _body(doc, "RSVP deadline: Friday, November 6, 2026.",
                     NAVY if fixed else GOLD)  # Defect when gold: 3.64:1 on white.
    deadline.runs[0].bold = True

    doc.save(str(path))
    _postprocess(path, strip_alt=not fixed)


def main():
    (HERE / "assets").mkdir(exist_ok=True)
    quad_header(HERE / "assets" / "quad_header.png")
    rsvp_chart(HERE / "assets" / "rsvp_chart.png")
    build_memo(HERE / "memo_before.docx", fixed=False)
    build_memo(HERE / "memo_after.docx", fixed=True)
    print("built sample/assets/*.png, sample/memo_before.docx, sample/memo_after.docx")


if __name__ == "__main__":
    main()
